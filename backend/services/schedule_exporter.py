"""Exporta l'horari actiu (o una proposta) a un únic fitxer Excel amb una
pestanya per a cada grup, professor i aula.

No depèn de cap repositori de calendari: dedueix l'ordre dels dies a partir
d'una llista coneguda de noms en català, i l'ordre de les hores parsejant
directament el text ("8:00", "8:30"...). Si algun dia/hora no coincideix amb
el format esperat, simplement es col·loca al final, ordenat alfabèticament,
en lloc de fallar.
"""

from __future__ import annotations

import re
from io import BytesIO
from pathlib import Path
from typing import Any, Dict, List, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

try:
    from openpyxl.drawing.image import Image as XLImage
except ImportError:  # pragma: no cover - Pillow no instal·lat
    XLImage = None

_LOGO_PATH = Path(__file__).resolve().parent.parent / "assets" / "logo-emad.png"
_LOGO_HEADER_ROWS = 4
_LOGO_TARGET_WIDTH_PX = 150

def _split_group_names(group_text: str) -> List[str]:
    """Separa un camp de grup que pot contenir diversos grups per coma
    (p.ex. 'GI, GP'), preservant el text original (majúscules incloses)
    per mostrar-lo bé al nom de la pestanya."""
    if not group_text:
        return []
    parts = [part.strip() for part in str(group_text).split(",")]
    unique: List[str] = []
    seen = set()
    for part in parts:
        key = part.casefold()
        if part and key not in seen:
            seen.add(key)
            unique.append(part)
    return unique

_DAY_ORDER = ["dilluns", "dimarts", "dimecres", "dijous", "divendres", "dissabte", "diumenge"]

_HEADER_FILL = PatternFill(start_color="FF2F3B52", end_color="FF2F3B52", fill_type="solid")
_HEADER_FONT = Font(color="FFFFFFFF", bold=True)
_CELL_FONT = Font(size=10)
_CELL_BORDER = Border(*(Side(style="thin", color="FFD9D9D9"),) * 4)
_WRAP = Alignment(wrap_text=True, vertical="top", horizontal="left")
_CENTER = Alignment(horizontal="center", vertical="center")


def _day_sort_key(day_name: str):
    normalized = str(day_name or "").strip().lower()
    try:
        return (0, _DAY_ORDER.index(normalized))
    except ValueError:
        return (1, normalized)


def _hour_sort_key(hour_label: str):
    match = re.match(r"\s*(\d+):(\d+)", str(hour_label or ""))
    if match:
        return (0, int(match.group(1)) * 60 + int(match.group(2)))
    return (1, str(hour_label or ""))


_HEADER_LABELS = {"group": "Grup", "teacher": "Professor", "room": "Aula"}
_TITLE_FONT = Font(size=14, bold=True, color="FF2F3B52")
_NOTE_FONT = Font(size=10, italic=True, color="FF555555")
_TITLE_ALIGNMENT = Alignment(horizontal="right", vertical="center")
_NOTE_ALIGNMENT = Alignment(horizontal="right", vertical="center", wrap_text=True)


def _insert_logo(sheet) -> None:
    """Insereix el logotip d'EMAD a la cantonada superior esquerra del full,
    si el fitxer existeix. Si no hi és (p.ex. en un entorn que no l'ha
    desplegat), simplement s'omet sense fer fallar l'exportació."""
    if not _LOGO_PATH.exists() or XLImage is None:
        return
    try:
        image = XLImage(str(_LOGO_PATH))
        original_width, original_height = image.width, image.height
        if original_width:
            image.width = _LOGO_TARGET_WIDTH_PX
            image.height = round(_LOGO_TARGET_WIDTH_PX * original_height / original_width)
        sheet.add_image(image, "A1")
    except Exception:
        return


def _write_title(sheet, sheet_kind: str, display_name: str, notes: List[str], last_column: int) -> None:
    """Escriu, alineat a la dreta i a l'alçada del logo (cantonada superior
    esquerra), el nom del grup/professor/aula d'aquest full. Si hi ha notes
    (p.ex. hores de Tutoria per a un grup), es mostren just a sota, també
    alineades a la dreta."""
    last_column = max(last_column, 2)
    label = _HEADER_LABELS.get(sheet_kind, "")
    title_text = f"{label}: {display_name}" if label else display_name

    title_row = 2
    sheet.merge_cells(start_row=title_row, start_column=2, end_row=title_row, end_column=last_column)
    title_cell = sheet.cell(row=title_row, column=2, value=title_text)
    title_cell.font = _TITLE_FONT
    title_cell.alignment = _TITLE_ALIGNMENT

    if notes:
        note_row = title_row + 1
        sheet.merge_cells(start_row=note_row, start_column=2, end_row=note_row, end_column=last_column)
        note_cell = sheet.cell(row=note_row, column=2, value="\n".join(notes))
        note_cell.font = _NOTE_FONT
        note_cell.alignment = _NOTE_ALIGNMENT


def _cell_text(activity: Dict[str, Any], sheet_kind: str = "group") -> str:
    parts = [activity.get("subject") or ""]
    if sheet_kind == "room":
        extra_fields = [activity.get("group"), activity.get("teacher")]
    elif sheet_kind == "teacher":
        extra_fields = [activity.get("group"), activity.get("room")]
    else:
        extra_fields = [activity.get("teacher"), activity.get("room")]
    extra = " · ".join(part for part in extra_fields if part)
    if extra:
        parts.append(extra)
    return "\n".join(parts)


def _is_tutoria(activity: Dict[str, Any]) -> bool:
    return (activity.get("subject") or "").strip().casefold() == "tutoria"


def _tutoria_note(activity: Dict[str, Any]) -> str:
    teacher = (activity.get("teacher") or "").strip()
    day = (activity.get("day") or "").strip()
    start = (activity.get("start") or "").strip()
    when = " ".join(part for part in [day, start] if part)
    detail = " — ".join(part for part in [teacher, when] if part)
    return f"Tutoria: {detail}" if detail else "Tutoria"


def _collect_axes(activities: Sequence[Dict[str, Any]]) -> tuple[List[str], List[str]]:
    days = sorted({a.get("day") for a in activities if a.get("day")}, key=_day_sort_key)
    hours = sorted({a.get("start") for a in activities if a.get("start")}, key=_hour_sort_key)
    return days, hours


def _write_grid_sheet(
    workbook: Workbook,
    title: str,
    activities: List[Dict[str, Any]],
    sheet_kind: str = "group",
    display_name: str = "",
    notes: List[str] = None,
) -> None:
    # Els noms de fulla d'Excel no poden superar 31 caràcters ni contenir
    # certs símbols; es netegen per evitar que openpyxl exploti.
    safe_title = re.sub(r"[\\/*?:\[\]]", " ", title).strip()[:31] or "Horari"
    base_title = safe_title
    suffix = 1
    existing_titles = {sheet.title for sheet in workbook.worksheets}
    while safe_title in existing_titles:
        suffix += 1
        safe_title = f"{base_title[:28]} ({suffix})"

    sheet = workbook.create_sheet(title=safe_title)
    _insert_logo(sheet)
    for row_index in range(1, _LOGO_HEADER_ROWS + 1):
        sheet.row_dimensions[row_index].height = 18

    header_row = _LOGO_HEADER_ROWS + 1

    days, hours = _collect_axes(activities)
    last_column = max(len(days) + 1, 2)
    _write_title(sheet, sheet_kind, display_name or title, notes or [], last_column)

    if not days or not hours:
        sheet.cell(row=header_row, column=1, value="Sense activitats programades.")
        return

    sheet.cell(row=header_row, column=1, value="Hora")
    for col_index, day in enumerate(days, start=2):
        cell = sheet.cell(row=header_row, column=col_index, value=day.capitalize())
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = _CENTER
    sheet.cell(row=header_row, column=1).fill = _HEADER_FILL
    sheet.cell(row=header_row, column=1).font = _HEADER_FONT
    sheet.cell(row=header_row, column=1).alignment = _CENTER

    grid: Dict[tuple, List[str]] = {}
    for activity in activities:
        key = (activity.get("day"), activity.get("start"))
        grid.setdefault(key, []).append(_cell_text(activity, sheet_kind))

    for row_offset, hour in enumerate(hours):
        row_index = header_row + 1 + row_offset
        hour_cell = sheet.cell(row=row_index, column=1, value=hour)
        hour_cell.font = Font(bold=True, size=10)
        hour_cell.alignment = _CENTER
        hour_cell.border = _CELL_BORDER

        for col_index, day in enumerate(days, start=2):
            texts = grid.get((day, hour), [])
            cell = sheet.cell(row=row_index, column=col_index, value="\n\n".join(texts) if texts else "")
            cell.font = _CELL_FONT
            cell.alignment = _WRAP
            cell.border = _CELL_BORDER

    sheet.column_dimensions["A"].width = 9
    for col_index in range(2, len(days) + 2):
        sheet.column_dimensions[get_column_letter(col_index)].width = 26
    for row_offset in range(len(hours)):
        sheet.row_dimensions[header_row + 1 + row_offset].height = 42
    sheet.freeze_panes = f"B{header_row + 1}"


def build_schedule_export(activities: Sequence[Dict[str, Any]]) -> BytesIO:
    """Retorna un .xlsx (com a BytesIO) amb una pestanya per a cada grup
    pare, professor i aula que apareguin a `activities`. Els grups
    combinats (p.ex. 'GI, GP') generen una pestanya per a cada grup
    individual, amb les mateixes activitats repetides a totes dues.

    Les activitats de "Tutoria" són un cas especial: no ocupen cap franja
    a l'horari del grup tutoritzat (només hi surten com una nota sota el
    nom del grup, amb el professor i l'hora), però sí que ocupen la seva
    franja normal a l'horari del professor, indicant el grup."""
    workbook = Workbook()
    workbook.remove(workbook.active)

    by_group: Dict[str, List[Dict[str, Any]]] = {}
    by_teacher: Dict[str, List[Dict[str, Any]]] = {}
    by_room: Dict[str, List[Dict[str, Any]]] = {}
    group_notes: Dict[str, List[str]] = {}

    for activity in activities:
        is_tutoria = _is_tutoria(activity)
        group_field_names = _split_group_names(activity.get("group") or "")

        if is_tutoria:
            for name in group_field_names:
                group_notes.setdefault(name, []).append(_tutoria_note(activity))
        else:
            for name in group_field_names:
                by_group.setdefault(name, []).append(activity)

        teacher_text = (activity.get("teacher") or "").strip()
        for teacher in [part.strip() for part in teacher_text.split(",") if part.strip()]:
            by_teacher.setdefault(teacher, []).append(activity)

        room_text = (activity.get("room") or "").strip()
        if room_text:
            by_room.setdefault(room_text, []).append(activity)

    all_group_names = set(by_group) | set(group_notes)
    for group_name in sorted(all_group_names):
        _write_grid_sheet(
            workbook,
            f"G-{group_name}",
            by_group.get(group_name, []),
            sheet_kind="group",
            display_name=group_name,
            notes=group_notes.get(group_name, []),
        )
    for teacher_name in sorted(by_teacher):
        _write_grid_sheet(
            workbook, f"P-{teacher_name}", by_teacher[teacher_name], sheet_kind="teacher", display_name=teacher_name
        )
    for room_name in sorted(by_room):
        _write_grid_sheet(
            workbook, f"A-{room_name}", by_room[room_name], sheet_kind="room", display_name=room_name
        )

    if not workbook.worksheets:
        sheet = workbook.create_sheet(title="Horari")
        _insert_logo(sheet)
        sheet.cell(row=_LOGO_HEADER_ROWS + 1, column=1, value="No hi ha cap activitat programada.")

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer

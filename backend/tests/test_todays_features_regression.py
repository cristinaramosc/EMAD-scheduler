"""Tests de regressió per a les funcionalitats afegides en la sessió de
treball sobre bloqueig d'horari de grup, Tutoria a les exportacions, i
exportació en PDF.

Cadascun d'aquests tests reprodueix el comportament esperat validat
manualment durant el desenvolupament; si algú torna a introduir un bug
relacionat, aquests tests han de fallar per avisar-ho de seguida.
"""

from io import BytesIO

from application.scheduler_use_cases import SchedulerUseCases
from repositories.academic_data_repository import AcademicDataRepository
from scheduler_engine.models import Activity
from services.schedule_exporter import build_schedule_export
from services.schedule_pdf_exporter import build_schedule_pdf


# ---------------------------------------------------------------------------
# Bloqueig d'horari de grup
# ---------------------------------------------------------------------------


class _FakeSnapshot:
    def __init__(self, active_schedule):
        self.active_schedule = active_schedule


def _make_use_cases(active_schedule):
    uc = SchedulerUseCases.__new__(SchedulerUseCases)
    uc._working_timetable_repo = object()  # només cal que no sigui None
    uc._load_snapshot = lambda: _FakeSnapshot(active_schedule)
    return uc


def test_locked_group_position_is_found_even_with_combined_groups():
    """Un grup bloquejat ('GI') dins d'un bloc combinat ('GI, GP') a
    l'horari actiu ha de trobar-se correctament."""
    uc = _make_use_cases(
        [
            {
                "id": 1,
                "teacher": "Bego",
                "subject": "Ll. Tec. Audio UF3",
                "group": "GI, GP",
                "room": "A1",
                "day": "Dimarts",
                "start": "10:00",
                "duration": 2,
            }
        ]
    )
    placements = uc._build_locked_placements({"gi"})
    assert placements[("gi", "ll. tec. audio uf3", "bego")] == ("Dimarts", "10:00")


def test_apply_locked_placement_injects_fixed_day_and_start_on_a_copy():
    uc = _make_use_cases([])
    placements = {("gi", "ll. tec. audio uf3", "bego"): ("Dimarts", "10:00")}
    assignment = {"group": "GI, GP", "subject": "Ll. Tec. Audio UF3", "teacher": "Bego", "weekly_hours": 2}

    result = uc._apply_locked_placement(assignment, {"gi"}, placements)

    assert result is not assignment, "ha de retornar una còpia, no mutar l'original"
    assert assignment.get("fixed_day") is None, "l'assignació original no s'ha de tocar"
    assert result["fixed_day"] == "Dimarts"
    assert result["fixed_start"] == "10:00"


def test_apply_locked_placement_leaves_unlocked_assignments_untouched():
    uc = _make_use_cases([])
    assignment = {"group": "1r COM", "subject": "FOL", "teacher": "Carme"}

    result = uc._apply_locked_placement(assignment, {"gi"}, {})

    assert result is assignment


def test_build_locked_placements_is_empty_without_locked_groups():
    uc = _make_use_cases([{"id": 1, "group": "GI", "subject": "X", "teacher": "Y", "day": "Dilluns", "start": "8:00"}])
    assert uc._build_locked_placements(set()) == {}


def test_quarter_pair_alignment_reuses_start_even_with_different_durations():
    uc = _make_use_cases([])
    uc._academic_data_repo = AcademicDataRepository()
    uc._time_labels = {
        "day_names": ["Dilluns"],
        "hour_names": ["8:00", "8:30", "9:00", "9:30"],
    }
    uc._school_calendar = type("Calendar", (), {
        "days": [0],
        "periods_per_day": 4,
        "periods_for_day": lambda self, day: [
            type("Slot", (), {"day": day, "period": period})()
            for period in range(4)
        ],
    })()
    uc._scheduler_engine = __import__("scheduler_engine.engine", fromlist=["SchedulerEngine"]).SchedulerEngine()

    proposal = type("Proposal", (), {
        "id": "quarters",
        "activities": [
            Activity(1, "Ana", "Projecte 1Q", "GI", "", "Dilluns", "8:00", 2),
            Activity(2, "Biel", "Projecte 2Q", "GI", "", "Dilluns", "9:00", 1),
        ],
        "score": 0,
        "warnings": [],
        "conflicts": [],
        "score_breakdown": None,
        "metadata": {},
    })()

    aligned = uc._apply_quarter_pair_alignment(proposal)

    assert aligned.activities[0].start == "9:00"
    assert aligned.activities[1].start == "9:00"


def test_quarter_pair_alignment_only_matches_opposite_subject_endings():
    uc = _make_use_cases([])
    uc._academic_data_repo = AcademicDataRepository()
    uc._time_labels = {"day_names": ["Dilluns"], "hour_names": ["8:00", "8:30", "9:00"]}
    uc._school_calendar = type("Calendar", (), {
        "days": [0],
        "periods_per_day": 3,
        "periods_for_day": lambda self, day: [
            type("Slot", (), {"day": day, "period": period})() for period in range(3)
        ],
    })()
    uc._scheduler_engine = __import__("scheduler_engine.engine", fromlist=["SchedulerEngine"]).SchedulerEngine()

    def aligned_start(first_subject, second_subject):
        proposal = type("Proposal", (), {
            "id": "quarters",
            "activities": [
                Activity(1, "Ana", first_subject, "GI", "", "Dilluns", "8:00", 1),
                Activity(2, "Biel", second_subject, "GI", "", "Dilluns", "9:00", 1),
            ],
            "score": 0,
            "warnings": [],
            "conflicts": [],
            "score_breakdown": None,
            "metadata": {},
        })()
        result = uc._apply_quarter_pair_alignment(proposal)
        return result.activities[0].start, result.activities[1].start

    assert aligned_start("Projecte 1Q", "Projecte 2Q") == ("9:00", "9:00")
    assert aligned_start("Projecte 1Q", "Projecte 1Q") == ("8:00", "9:00")
    assert aligned_start("Projecte 2Q", "Projecte 2Q") == ("8:00", "9:00")
    assert aligned_start("Projecte", "Projecte 2Q") == ("8:00", "9:00")




# ---------------------------------------------------------------------------
# Tutoria a l'exportació Excel: informativa al grup, bloc real al professor
# ---------------------------------------------------------------------------

_TUTORIA_ACTIVITIES = [
    {
        "id": 1,
        "teacher": "Carme",
        "subject": "FOL",
        "group": "1r COM",
        "room": "A2",
        "day": "Dilluns",
        "start": "10:00",
        "duration": 4,
    },
    {
        "id": 2,
        "teacher": "Marta",
        "subject": "Tutoria",
        "group": "1r APGI",
        "room": "",
        "day": "Dimarts",
        "start": "12:00",
        "duration": 2,
    },
]


def test_tutoria_does_not_occupy_a_grid_slot_in_the_group_sheet():
    from openpyxl import load_workbook

    buffer = build_schedule_export(_TUTORIA_ACTIVITIES)
    workbook = load_workbook(BytesIO(buffer.read()))

    assert "G-1r APGI" in workbook.sheetnames
    sheet = workbook["G-1r APGI"]
    # No hi ha d'haver cap altra activitat que Tutoria per a aquest grup, i
    # per tant la graella ha de quedar buida (només el missatge "Sense
    # activitats programades.").
    found_grid_text = False
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value and "Tutoria" in str(cell.value) and "—" not in str(cell.value):
                found_grid_text = True
    assert not found_grid_text, "la Tutoria no ha d'aparèixer com a bloc a la graella del grup"


def test_tutoria_appears_as_a_note_under_the_group_title():
    from openpyxl import load_workbook

    buffer = build_schedule_export(_TUTORIA_ACTIVITIES)
    workbook = load_workbook(BytesIO(buffer.read()))
    sheet = workbook["G-1r APGI"]

    note_texts = [str(cell.value) for row in sheet.iter_rows() for cell in row if cell.value]
    assert any("Tutoria: Marta" in text and "Dimarts" in text for text in note_texts)


def test_tutoria_appears_as_a_real_block_on_the_teacher_sheet():
    from openpyxl import load_workbook

    buffer = build_schedule_export(_TUTORIA_ACTIVITIES)
    workbook = load_workbook(BytesIO(buffer.read()))

    assert "P-Marta" in workbook.sheetnames
    sheet = workbook["P-Marta"]
    cell_texts = [str(cell.value) for row in sheet.iter_rows() for cell in row if cell.value]
    assert any("Tutoria" in text and "1r APGI" in text for text in cell_texts)


# ---------------------------------------------------------------------------
# Exportació en PDF: prova de fum
# ---------------------------------------------------------------------------


def test_pdf_export_builds_one_page_per_group_teacher_and_room():
    buffer = build_schedule_pdf(_TUTORIA_ACTIVITIES)
    data = buffer.read()

    assert data[:5] == b"%PDF-"
    real_page_count = data.count(b"/Type /Page") - data.count(b"/Type /Pages")
    # 2 grups (1r COM, 1r APGI) + 2 professors (Carme, Marta) + 1 aula (A2)
    assert real_page_count == 5


def test_pdf_export_handles_empty_schedule_without_crashing():
    buffer = build_schedule_pdf([])
    assert buffer.getbuffer().nbytes > 0


# ---------------------------------------------------------------------------
# Compactació d'activitats: grups combinats han de respectar les franges
# no disponibles de cadascun dels seus noms individuals
# ---------------------------------------------------------------------------


class _FakeAcademicDataRepoForCompaction:
    """Només exposa el mínim que `_collect_blocked_slots` necessita."""

    def __init__(self, group_restrictions):
        self._group_restrictions = group_restrictions

    def active_teacher_restrictions(self):
        return []

    def active_group_restrictions(self):
        return self._group_restrictions


def _make_use_cases_for_compaction(group_restrictions, hour_names=None):
    uc = SchedulerUseCases.__new__(SchedulerUseCases)
    uc._time_labels = {
        "day_names": ["Dilluns", "Dimarts", "Dimecres", "Dijous", "Divendres"],
        "hour_names": hour_names or ["8:00", "8:30", "9:00", "10:00", "10:30", "11:00", "15:00", "15:30"],
    }
    uc._academic_data_repo = _FakeAcademicDataRepoForCompaction(group_restrictions)
    return uc


def test_combined_group_compaction_respects_each_individual_groups_blocked_slots():
    """Bug real: una activitat de grup combinat ('GI, GP') es compactava
    fins a les 8:00 del matí perquè `_compact_activities` només mirava el
    text exacte 'GI, GP' a les franges bloquejades, que mai coincidia amb
    les restriccions reals (guardades per 'GI' i 'GP' per separat). Ha de
    respectar la unió de les franges bloquejades de tots dos noms."""
    uc = _make_use_cases_for_compaction(
        [
            {"group": "GI", "unavailable_slots": ["Dilluns 8:00", "Dilluns 8:30", "Dilluns 9:00"]},
            {"group": "GP", "unavailable_slots": ["Dilluns 8:00", "Dilluns 8:30", "Dilluns 9:00"]},
        ]
    )
    activities = [
        Activity(
            id=1,
            teacher="Noëlle",
            subject="Foto UF3",
            group="GI, GP",
            room="",
            day="Dilluns",
            start="15:00",
            duration=4,
        )
    ]

    compacted, moved_ids = uc._compact_activities(activities)

    assert compacted[0].start == "15:00"
    assert moved_ids == []


def test_single_group_compaction_still_works_as_before():
    """No-regressió: una activitat de grup individual (no combinat) es
    continua compactant normalment cap a la primera franja lliure."""
    uc = _make_use_cases_for_compaction(
        [{"group": "GI", "unavailable_slots": ["Dilluns 8:00"]}]
    )
    activities = [
        Activity(
            id=1,
            teacher="Marc",
            subject="GPP",
            group="GI",
            room="",
            day="Dilluns",
            start="9:00",
            duration=2,
        )
    ]

    compacted, moved_ids = uc._compact_activities(activities)

    assert compacted[0].start == "8:30"
    assert moved_ids == [1]


def test_group_compaction_respects_group_daily_start_time():
    """La compactació no ha de forçar el grup a començar a les 8:00 si la
    restricció de grup exigeix començar més tard (p.ex. 10:00)."""
    uc = _make_use_cases_for_compaction(
        [{"group": "GI", "daily_start_time": "10:00"}],
        hour_names=["8:00", "8:30", "9:00", "9:30", "10:00", "10:30", "11:00", "11:30", "15:00", "15:30"],
    )
    activities = [
        Activity(
            id=1,
            teacher="Marc",
            subject="GPP",
            group="GI",
            room="",
            day="Dilluns",
            start="15:00",
            duration=2,
        ),
        Activity(
            id=2,
            teacher="Marc",
            subject="GPP 2",
            group="GI",
            room="",
            day="Dilluns",
            start="15:30",
            duration=1,
        ),
    ]

    compacted, moved_ids = uc._compact_activities(activities)

    assert [activity.start for activity in compacted] == ["10:00", "11:00"]
    assert sorted(moved_ids) == [1, 2]


def test_compaction_merges_transitive_combined_group_components():
    uc = _make_use_cases_for_compaction([])
    activities = [
        Activity(1, "A", "GI", "GI", "", "Dilluns", "8:00", 1),
        Activity(2, "B", "GP", "GP", "", "Dilluns", "10:00", 1),
        Activity(3, "C", "Compartida", "GI, GP", "", "Dilluns", "15:00", 1),
    ]

    compacted, _ = uc._compact_activities(activities)

    assert [activity.start for activity in compacted] == ["8:00", "8:30", "9:00"]
